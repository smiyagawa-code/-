#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""デモで落とす見本ファイル（すべて架空）を作る。

  python3 scripts/make-samples.py            # public/sample/ と ~/Downloads/⭐️ACS様デモ_2026-10-06/ に出す
  python3 scripts/make-samples.py --no-copy  # public/sample/ だけ

作るもの（public/sample/）
  内示_今回_2026-09-25.xlsx       CSV と同じ内容を Excel に（openpyxl）
  内示_今回_2026-09-25.pdf        得意先が送ってくる内示書らしい PDF 2 ページ（reportlab。日本語フォントを埋め込む）
  内示_メール本文_2026-09-25.txt  得意先担当者 → 購買部 高橋 宛のメール本文
  内示_電話メモ_2026-09-25.txt    電話で聞いた内容の走り書き（表記ゆれあり）
  内示_第3工場_初回_2026-09-30.txt 新しい案件（第3工場 検査ライン新設）の初回の電話メモ → 案件を新しく作る用
  内示_第3工場_2026-10-03.pdf     同じ新しい案件の内示書 PDF 1 ページ（3台・前倒し）→ 作った案件に入れる用
  内示_第3工場_2026-10-03.csv     上の PDF と同じ内容の CSV（列順・BOM 無しは 内示_今回 CSV と同じ）
  はじめにお読みください.txt      Downloads 側にだけ置く（デモの落とす順）

  python3 scripts/make-samples.py --new-only  # 新しい案件の 3 本だけ作って Downloads に足す（他は消さない）

数字・文言は src/data.js の FOLDERS（株式会社大和精密製作所の 3 案件）と
public/sample/内示_今回_2026-09-25.csv に合わせる。変えるときは両方を直すこと。

必要なもの: python3 / openpyxl / reportlab（無ければその旨を出して、作れる分だけ作る）
日本語フォント: macOS の BIZ UDGothic（TrueType）→ 無ければ Osaka → Arial Unicode の順で探す
"""
import glob
import os
import re
import shutil
import subprocess
import sys
from datetime import date

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLE = os.path.join(ROOT, 'public', 'sample')
DOWNLOADS = os.path.expanduser('~/Downloads/⭐️ACS様デモ_2026-10-06')

# ---------- 架空データ（src/data.js の FOLDERS と同じ） ----------
CUSTOMER = '株式会社大和精密製作所'
VENDOR = 'ACS株式会社'
ISSUE_DATE = '2026-09-25'
PREV_DATE = '2026-08-28'
DOC_NO = 'FC-2609-0131'            # 架空の内示番号
SENDER = {'dept': '生産管理部 生産計画課', 'name': '田村 由紀', 'tel': '042-000-0000（内線 318）', 'mail': 'y.tamura@example.co.jp'}
RECEIVER = {'dept': '購買部', 'name': '高橋'}

# 列: 案件, 機種, 前回数量, 今回数量, 前回希望日, 今回希望日, 前回仕様, 今回仕様, 備考
CASES = [
    {
        'name': '第2工場 外観検査ライン増設', 'model': 'VIS-200',
        'qty_prev': 2, 'qty': 3, 'due_prev': '2026-11-14', 'due': '2026-11-14',
        'opt_prev': '', 'opt': 'NG排出シュート',
        'note': '3号機は増産対応のため追加。NG排出シュートは3台とも。できれば 11/7',
        'hand': 'できれば 11/7',   # PDF では手書きふうに添える
        'change': '台数 2台 → 3台（3号機を追加）。NG排出シュートを 3台とも',
    },
    {
        'name': '組立セル AS-500 導入', 'model': 'AS-500',
        'qty_prev': 2, 'qty': 2, 'due_prev': '2026-11-10', 'due': '2026-10-15',
        'opt_prev': '安全柵', 'opt': 'ライトカーテン仕様',
        'note': '工場レイアウト変更に伴い前倒し。安全柵はライトカーテン仕様に変更',
        'hand': '',
        'change': '希望日 11/10 → 10/15 に前倒し。安全柵 → ライトカーテン仕様',
    },
    {
        'name': 'クリーン仕様 搬送ユニット', 'model': 'CV-CL',
        'qty_prev': 1, 'qty': 0, 'due_prev': '2026-12-05', 'due': '2026-12-05',
        'opt_prev': '', 'opt': '',
        'note': 'クリーンルーム計画の見直しのため取消。再開時期は未定',
        'hand': '',
        'change': '1台 → 取消（0台）',
    },
]
HEADER = ['内示日', '得意先', '案件', '機種', '数量', '納入希望日', '仕様', '備考']

# 新しい案件（デモで「案件を作る → 内示を入れて変化を見せる」用）。機種は第2工場と同じ VIS-200 なので部品表を流用できる
NEW_CASE = {
    'name': '第3工場 検査ライン新設', 'model': 'VIS-200',
    'doc_no': 'FC-2610-0007',                     # 架空の内示番号
    # 初回: 9/30 の電話（内示書はまだ無い）
    'first': {'date': '2026-09-30', 'qty': 2, 'due': '2026-12-10', 'opt': '', 'note': '予算承認は 10 月中'},
    # 内示書: 10/3 に PDF と CSV で届く
    'now': {'date': '2026-10-03', 'qty': 3, 'due': '2026-11-28', 'opt': 'NG排出シュート', 'note': 'ライン立ち上げを前倒し。3台目は予備機'},
}


def jp(d):
    """2026-09-25 → 2026年9月25日"""
    y, m, dd = d.split('-')
    return f'{y}年{int(m)}月{int(dd)}日'


def md(d):
    y, m, dd = d.split('-')
    return f'{int(m)}/{int(dd)}'


def ok(msg):
    print('  ' + msg)


# ---------- 1. Excel ----------
def make_xlsx(path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print('  [できません] openpyxl が入っていないので Excel は作れません（pip は使いません）')
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = '内示'
    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill('solid', fgColor='1F4E78')
    head_font = Font(name='BIZ UDPゴシック', bold=True, color='FFFFFF', size=11)
    body_font = Font(name='BIZ UDPゴシック', size=11)
    zebra = PatternFill('solid', fgColor='EEF3F8')

    # 1 行目が見出し（読み取り側が先頭行を見出しとして扱えるように、タイトル行は別シートに置く）
    ws.append(HEADER)
    for c in ws[1]:
        c.fill, c.font, c.border = head_fill, head_font, border
        c.alignment = Alignment(horizontal='center', vertical='center')
    for i, cs in enumerate(CASES, start=2):
        ws.append([ISSUE_DATE, CUSTOMER, cs['name'], cs['model'], cs['qty'], cs['due'], cs['opt'], cs['note']])
        for c in ws[i]:
            c.font, c.border = body_font, border
            c.alignment = Alignment(vertical='center', wrap_text=True)
            if i % 2 == 1:
                c.fill = zebra
        ws.cell(row=i, column=5).alignment = Alignment(horizontal='right', vertical='center')
        ws.cell(row=i, column=5).number_format = '0'
        ws.cell(row=i, column=5).font = Font(name='BIZ UDPゴシック', size=11, bold=True, color='C00000' if cs['qty'] == 0 else '000000')
    widths = [12, 26, 30, 10, 7, 14, 20, 56]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22
    for i in range(2, 2 + len(CASES)):
        ws.row_dimensions[i].height = 34
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADER))}{1 + len(CASES)}'
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1

    # 送付情報は別シート
    cover = wb.create_sheet('送付状')
    rows = [
        ['内示書（フォーキャスト）'],
        [],
        ['宛先', f'{VENDOR} {RECEIVER["dept"]} {RECEIVER["name"]} 様'],
        ['発行', f'{CUSTOMER} {SENDER["dept"]} {SENDER["name"]}'],
        ['発行日', jp(ISSUE_DATE)],
        ['内示番号', DOC_NO],
        ['前回', f'{jp(PREV_DATE)} 発行分からの変更です'],
        [],
        ['ご注意', '本内示は現時点の生産計画にもとづく見込みであり、注文を確約するものではありません。'],
        ['', '数量・納期は変更になる場合があります。変更の際はあらためてご連絡します。'],
        [],
        ['※', 'このファイルは営業デモ用の架空データです。実在の会社・人物・取引とは関係ありません。'],
    ]
    for r in rows:
        cover.append(r)
    cover['A1'].font = Font(name='BIZ UDPゴシック', size=16, bold=True, color='1F4E78')
    for r in range(3, 13):
        cover.cell(row=r, column=1).font = Font(name='BIZ UDPゴシック', size=11, bold=True)
        cover.cell(row=r, column=2).font = body_font
    cover.column_dimensions['A'].width = 12
    cover.column_dimensions['B'].width = 80
    cover.sheet_view.showGridLines = False

    wb.save(path)
    return True


# ---------- 2. PDF ----------
def find_font():
    """TrueType アウトラインの日本語フォント（reportlab は CFF のヒラギノを読めない）"""
    cands = [
        ('BIZ_UDGothic', glob.glob('/System/Library/AssetsV2/com_apple_MobileAsset_Font*/*/AssetData/BIZ_UDGothic.ttc'), 0, 1),
        ('Osaka', glob.glob('/System/Library/AssetsV2/com_apple_MobileAsset_Font*/*/AssetData/Osaka.ttf'), 0, 0),
        ('ArialUnicode', ['/Library/Fonts/Arial Unicode.ttf', '/System/Library/Fonts/Supplemental/Arial Unicode.ttf'], 0, 0),
    ]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for name, paths, ri, bi in cands:
        for p in paths:
            if not os.path.exists(p):
                continue
            try:
                pdfmetrics.registerFont(TTFont('JP', p, subfontIndex=ri))
                pdfmetrics.registerFont(TTFont('JP-B', p, subfontIndex=bi))
                return name, p
            except Exception as e:  # noqa: BLE001
                print(f'  （{name} は使えません: {e}）')
    return None, None


def make_pdf(path):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_RIGHT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        print('  [できません] reportlab が入っていないので PDF は作れません（pip は使いません）')
        return False
    font_name, font_path = find_font()
    if not font_name:
        print('  [できません] 埋め込める日本語 TrueType フォントが見つからないので PDF は作れません')
        return False
    ok(f'フォント: {font_name}（{font_path}）')

    navy = colors.HexColor('#1F4E78')
    st = ParagraphStyle('b', fontName='JP', fontSize=10, leading=15)
    st_s = ParagraphStyle('s', parent=st, fontSize=8.5, leading=12, textColor=colors.HexColor('#444444'))
    st_r = ParagraphStyle('r', parent=st, alignment=TA_RIGHT)
    st_h1 = ParagraphStyle('h1', parent=st, fontName='JP-B', fontSize=18, leading=26, textColor=navy)
    st_h2 = ParagraphStyle('h2', parent=st, fontName='JP-B', fontSize=12, leading=18, textColor=navy, spaceBefore=6, spaceAfter=4)
    st_cell = ParagraphStyle('c', parent=st, fontSize=9, leading=13)
    st_cell_b = ParagraphStyle('cb', parent=st_cell, fontName='JP-B')
    st_hand = ParagraphStyle('hand', parent=st_cell, textColor=colors.HexColor('#1A4FBF'))  # 手書きのつもり（青ペン）
    st_foot = ParagraphStyle('f', parent=st, fontSize=8, textColor=colors.grey)

    def header_block(title):
        return [
            Paragraph(title, st_h1),
            Spacer(1, 2 * mm),
            Table([
                [Paragraph(f'{VENDOR} 御中<br/>{RECEIVER["dept"]} {RECEIVER["name"]} 様', st),
                 Paragraph(f'内示番号: {DOC_NO}<br/>発行日: {jp(ISSUE_DATE)}<br/>{CUSTOMER}<br/>{SENDER["dept"]} {SENDER["name"]}<br/>TEL {SENDER["tel"]}', st_r)],
            ], colWidths=[95 * mm, 85 * mm], style=TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP')])),
            Spacer(1, 4 * mm),
        ]

    story = []
    # ---- 1 ページ目: 内示の一覧 ----
    story += header_block('内示書（フォーキャスト）')
    story.append(Paragraph(
        '平素よりお世話になっております。下記のとおり、当社の生産計画にもとづく内示をご連絡いたします。'
        f'{jp(PREV_DATE)} 発行分からの変更を含みますので、ご確認のうえ部品手配をお願いいたします。', st))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph('1. 内示内容', st_h2))
    st_head = ParagraphStyle('hd', parent=st_cell_b, textColor=colors.white)
    grid = [
        ('BACKGROUND', (0, 0), (-1, 0), navy),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#999999')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EEF3F8')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]
    head = ['No.', '案件名', '機種', '数量', '納入希望日', '仕様', '備考']
    rows = [[Paragraph(h, st_head) for h in head]]
    for i, cs in enumerate(CASES, start=1):
        note = cs['note'].replace('。できれば 11/7', '')
        note_p = Paragraph(note, st_cell)
        if cs['hand']:
            note_p = [Paragraph(note, st_cell), Paragraph(f'（手書き）{cs["hand"]}', st_hand)]
        qty = '0（取消）' if cs['qty'] == 0 else str(cs['qty'])
        rows.append([Paragraph(str(i), st_cell), Paragraph(cs['name'], st_cell), Paragraph(cs['model'], st_cell),
                     Paragraph(qty, st_cell_b), Paragraph(jp(cs['due']), st_cell),
                     Paragraph(cs['opt'] or '標準', st_cell), note_p])
    t = Table(rows, colWidths=[9 * mm, 39 * mm, 17 * mm, 20 * mm, 29 * mm, 26 * mm, 40 * mm], repeatRows=1)
    t.setStyle(TableStyle(grid))
    story.append(t)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph('2. ご注意', st_h2))
    for s in [
        '本内示は現時点の生産計画にもとづく見込みであり、注文を確約するものではありません。',
        '数量・納期は変更になる場合があります。変更の際はあらためてご連絡します。',
        '正式な注文書は、各案件の納入希望日の 4 週間前を目安に発行します。',
    ]:
        story.append(Paragraph('・' + s, st))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph('※ この文書は営業デモ用の架空データです。実在の会社・人物・取引とは関係ありません。', st_foot))

    # ---- 2 ページ目: 前回からの変更点 ----
    story.append(PageBreak())
    story += header_block('内示書（フォーキャスト）　別紙：前回からの変更点')
    story.append(Paragraph(f'3. {jp(PREV_DATE)} 発行分との差異', st_h2))
    head2 = ['No.', '案件名', '項目', '前回（8/28）', '今回（9/25）', '理由']
    rows2 = [[Paragraph(h, st_head) for h in head2]]
    diffs = [
        (1, CASES[0]['name'], '数量', '2', '3', '3号機は増産対応のため追加'),
        (1, CASES[0]['name'], '仕様', '標準', 'NG排出シュート（3台とも）', '歩留まり対策'),
        (2, CASES[1]['name'], '納入希望日', jp(CASES[1]['due_prev']), jp(CASES[1]['due']), '工場レイアウト変更に伴い前倒し'),
        (2, CASES[1]['name'], '仕様', '安全柵', 'ライトカーテン仕様', '安全柵はライトカーテン仕様に変更'),
        (3, CASES[2]['name'], '数量', '1', '0（取消）', 'クリーンルーム計画の見直しのため取消。再開時期は未定'),
    ]
    for d in diffs:
        rows2.append([Paragraph(str(d[0]), st_cell), Paragraph(d[1], st_cell), Paragraph(d[2], st_cell),
                      Paragraph(d[3], st_cell), Paragraph(d[4], st_cell_b), Paragraph(d[5], st_cell)])
    t2 = Table(rows2, colWidths=[9 * mm, 40 * mm, 22 * mm, 30 * mm, 36 * mm, 43 * mm], repeatRows=1)
    t2.setStyle(TableStyle(grid))
    story.append(t2)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph('4. 補足', st_h2))
    for s in [
        f'{CASES[0]["name"]}: 検査ラインの立ち上げ日程の都合で、可能であれば 11月7日 の納入をご検討ください（1 ページ目 備考欄の手書き）。',
        f'{CASES[1]["name"]}: ライトカーテンの段数は設備担当と確認中です。決まり次第ご連絡します。',
        f'{CASES[2]["name"]}: 取消は内示の取消です。すでに手配済みの部品があればお知らせください。',
    ]:
        story.append(Paragraph('・' + s, st))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph(f'お問い合わせ: {CUSTOMER} {SENDER["dept"]} {SENDER["name"]}（TEL {SENDER["tel"]} / {SENDER["mail"]}）', st_s))
    story.append(Paragraph('※ この文書は営業デモ用の架空データです。実在の会社・人物・取引とは関係ありません。', st_foot))

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont('JP', 8)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f'{DOC_NO}　{doc.page} / 2')
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=18 * mm,
                            title='内示書（フォーキャスト）', author=f'{CUSTOMER} {SENDER["dept"]}', subject='営業デモ用の架空データ')
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return True


# ---------- 3. メール本文 ----------
def make_mail(path):
    c1, c2, c3 = CASES
    body = f"""差出人: {SENDER['name']} <{SENDER['mail']}>
宛先: {VENDOR} {RECEIVER['dept']} {RECEIVER['name']} 様
日時: {jp(ISSUE_DATE)} 16:42
件名: 【内示】9月分の内示について（8/28発行分からの変更あり）

{VENDOR}
{RECEIVER['dept']} {RECEIVER['name']} 様

いつもお世話になっております。
{CUSTOMER} {SENDER['dept']}の{SENDER['name'].split()[0]}です。

9月分の内示をご連絡します。8月28日にお送りした内示から
3件とも変更がありますので、お手数ですがご確認ください。

1. {c1['name']}（{c1['model']}）
   台数を 2台から 3台に変更します。3号機は増産対応のための追加です。
   納入希望日は 11月14日のままですが、検査ラインの立ち上げ日程の都合で、
   可能であれば 11月7日に納めていただけると助かります。
   NG排出シュートは 3台とも付けてください。

2. {c2['name']}（{c2['model']}）
   台数は 2台で変わりません。
   工場レイアウトの変更に伴い、納入希望日を 11月10日から 10月15日に前倒しします。
   あわせて、安全柵はライトカーテン仕様に変更してください。
   （段数は設備担当と確認中です。決まり次第ご連絡します）

3. {c3['name']}（{c3['model']}）
   クリーンルーム計画の見直しのため、1台の内示を取消します。
   再開時期は未定です。すでに手配済みの部品があればお知らせください。

まとめると次のとおりです。
  ・{c1['model']}  2台 → 3台　希望日 11/14（できれば 11/7）　NG排出シュート 3台とも
  ・{c2['model']}  2台のまま　希望日 11/10 → 10/15　安全柵 → ライトカーテン仕様
  ・{c3['model']}  1台 → 取消

正式な内示書（PDF）は別途お送りします。
ご不明な点があればご連絡ください。
よろしくお願いいたします。

--
{CUSTOMER}
{SENDER['dept']}　{SENDER['name']}
TEL {SENDER['tel']}
{SENDER['mail']}

※ このメールは営業デモ用の架空データです。実在の会社・人物・取引とは関係ありません。
"""
    with open(path, 'w', encoding='utf-8') as f:
        f.write(body)
    return True


# ---------- 4. 電話メモ ----------
def make_memo(path):
    body = f"""9/25 15:10 大和精密 田村さんより TEL（高橋 受）

内示 9月分 3件とも変更あり とのこと

・VIS-200 外観検査 第2工場
  2台→3台に。3号機 増産対応で追加
  希望日 11/14 かわらず。でも「できれば11/7」と言ってた → 要確認
  NGシュート 3台ぜんぶ

・AS-500 2台 10/15に前倒し ライトカーテンに変更（安全柵やめる）
  レイアウト変更のため
  段数きいてない！ 4段？2段？ → 折り返し確認

・クリーン搬送 1台キャンセル
  クリーンルーム計画見直し、再開は未定
  手配ずみの部品あれば教えてほしいと

PDF の内示書はあとでメールで送ってくれるそう
→ 部品の影響 確認して 明日中に返事

※ 営業デモ用の架空メモです。実在の会社・人物・取引とは関係ありません。
"""
    with open(path, 'w', encoding='utf-8') as f:
        f.write(body)
    return True


# ---------- 4b. 新しい案件: 初回の電話メモ ----------
def make_new_memo(path):
    f = NEW_CASE['first']
    body = f"""{md(f['date'])} 10:40 大和精密 田村さんより TEL（高橋 受）

新規 第3工場に検査ライン新設 とのこと（内示書はまだ）

・{NEW_CASE['model']} {f['qty']}台 第3工場 検査ライン新設
  第2工場と同じ機種 → 部品表は第2工場のを流用でよさそう
  希望 {md(f['due'])}
  仕様は標準（NGシュートなし）

・予算承認は10月中 とのこと
  承認おりたら内示書くれる

→ 案件 新しく作っておく。部品の在庫 第2工場分とあわせて見る

※ 営業デモ用の架空メモです。実在の会社・人物・取引とは関係ありません。
"""
    with open(path, 'w', encoding='utf-8') as f_:
        f_.write(body)
    return True


# ---------- 4c. 新しい案件: 内示書 PDF（1 ページ） ----------
def make_new_pdf(path):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_RIGHT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        print('  [できません] reportlab が入っていないので PDF は作れません（pip は使いません）')
        return False
    font_name, font_path = find_font()
    if not font_name:
        print('  [できません] 埋め込める日本語 TrueType フォントが見つからないので PDF は作れません')
        return False
    ok(f'フォント: {font_name}（{font_path}）')

    first, now = NEW_CASE['first'], NEW_CASE['now']
    doc_no, issue = NEW_CASE['doc_no'], now['date']

    # 体裁は make_pdf と同じ
    navy = colors.HexColor('#1F4E78')
    st = ParagraphStyle('b', fontName='JP', fontSize=10, leading=15)
    st_s = ParagraphStyle('s', parent=st, fontSize=8.5, leading=12, textColor=colors.HexColor('#444444'))
    st_r = ParagraphStyle('r', parent=st, alignment=TA_RIGHT)
    st_h1 = ParagraphStyle('h1', parent=st, fontName='JP-B', fontSize=18, leading=26, textColor=navy)
    st_h2 = ParagraphStyle('h2', parent=st, fontName='JP-B', fontSize=12, leading=18, textColor=navy, spaceBefore=6, spaceAfter=4)
    st_cell = ParagraphStyle('c', parent=st, fontSize=9, leading=13)
    st_cell_b = ParagraphStyle('cb', parent=st_cell, fontName='JP-B')
    st_head = ParagraphStyle('hd', parent=st_cell_b, textColor=colors.white)
    st_foot = ParagraphStyle('f', parent=st, fontSize=8, textColor=colors.grey)
    grid = [
        ('BACKGROUND', (0, 0), (-1, 0), navy),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#999999')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EEF3F8')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]

    story = [
        Paragraph('内示書（フォーキャスト）', st_h1),
        Spacer(1, 2 * mm),
        Table([
            [Paragraph(f'{VENDOR} 御中<br/>{RECEIVER["dept"]} {RECEIVER["name"]} 様', st),
             Paragraph(f'内示番号: {doc_no}<br/>発行日: {jp(issue)}<br/>{CUSTOMER}<br/>{SENDER["dept"]} {SENDER["name"]}<br/>TEL {SENDER["tel"]}', st_r)],
        ], colWidths=[95 * mm, 85 * mm], style=TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP')])),
        Spacer(1, 4 * mm),
        Paragraph(
            '平素よりお世話になっております。下記のとおり、新規案件の内示をご連絡いたします。'
            f'{md(first["date"])} にお電話でお伝えした内容から変更がありますので、ご確認のうえ部品手配をお願いいたします。', st),
        Spacer(1, 3 * mm),
        Paragraph('1. 内示内容', st_h2),
    ]
    head = ['No.', '案件名', '機種', '数量', '納入希望日', '仕様', '備考']
    rows = [[Paragraph(h, st_head) for h in head],
            [Paragraph('1', st_cell), Paragraph(NEW_CASE['name'], st_cell), Paragraph(NEW_CASE['model'], st_cell),
             Paragraph(str(now['qty']), st_cell_b), Paragraph(jp(now['due']), st_cell),
             Paragraph(now['opt'] or '標準', st_cell), Paragraph(now['note'], st_cell)]]
    t = Table(rows, colWidths=[9 * mm, 39 * mm, 17 * mm, 20 * mm, 29 * mm, 26 * mm, 40 * mm], repeatRows=1)
    t.setStyle(TableStyle(grid))
    story += [t, Spacer(1, 4 * mm), Paragraph(f'2. {md(first["date"])} お電話でのご連絡からの変更点', st_h2)]
    head2 = ['項目', f'電話連絡（{md(first["date"])}）', f'今回（{md(issue)}）', '理由']
    diffs = [
        ('数量', str(first['qty']), str(now['qty']), '3台目は予備機'),
        ('納入希望日', jp(first['due']), jp(now['due']), 'ライン立ち上げを前倒し'),
        ('仕様', first['opt'] or '標準', now['opt'], '第2工場と同じ仕様にそろえる'),
    ]
    rows2 = [[Paragraph(h, st_head) for h in head2]]
    for d in diffs:
        rows2.append([Paragraph(d[0], st_cell), Paragraph(d[1], st_cell), Paragraph(d[2], st_cell_b), Paragraph(d[3], st_cell)])
    t2 = Table(rows2, colWidths=[28 * mm, 42 * mm, 50 * mm, 60 * mm], repeatRows=1)
    t2.setStyle(TableStyle(grid))
    story += [t2, Spacer(1, 4 * mm), Paragraph('3. ご注意', st_h2)]
    for s_ in [
        '本内示は現時点の生産計画にもとづく見込みであり、注文を確約するものではありません。',
        '数量・納期は変更になる場合があります。変更の際はあらためてご連絡します。',
        '正式な注文書は、納入希望日の 4 週間前を目安に発行します。',
        '予算承認は 10 月中の見込みです。承認後にあらためてご連絡します。',
    ]:
        story.append(Paragraph('・' + s_, st))
    story += [
        Spacer(1, 6 * mm),
        Paragraph(f'お問い合わせ: {CUSTOMER} {SENDER["dept"]} {SENDER["name"]}（TEL {SENDER["tel"]} / {SENDER["mail"]}）', st_s),
        Paragraph('※ この文書は営業デモ用の架空データです。実在の会社・人物・取引とは関係ありません。', st_foot),
    ]

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont('JP', 8)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f'{doc_no}　{doc.page} / 1')
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=18 * mm,
                            title='内示書（フォーキャスト）', author=f'{CUSTOMER} {SENDER["dept"]}', subject='営業デモ用の架空データ')
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return True


# ---------- 4d. 新しい案件: 内示書と同じ内容の CSV ----------
def make_new_csv(path):
    """内示_今回_2026-09-25.csv と同じ: 列順は HEADER、UTF-8（BOM 無し）、LF、末尾に改行"""
    import csv
    now = NEW_CASE['now']
    with open(path, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(HEADER)
        w.writerow([now['date'], CUSTOMER, NEW_CASE['name'], NEW_CASE['model'], now['qty'], now['due'], now['opt'], now['note']])
    return True


def check_new_pdf(path):
    """新しい案件の PDF: 1 ページで、主要な語が pdftotext で読めるか"""
    if not shutil.which('pdftotext'):
        ok('pdftotext が無いので文字化けの自動確認はしていません（PDF を開いて目で見てください）')
        return
    txt = subprocess.run(['pdftotext', path, '-'], capture_output=True, text=True).stdout
    flat = re.sub(r'\s+', '', txt)
    musts = [CUSTOMER, VENDOR, '御中', '内示書', NEW_CASE['name'], NEW_CASE['model'], 'NG排出シュート', '3台目は予備機', NEW_CASE['doc_no']]
    missing = [m for m in musts if re.sub(r'\s+', '', m) not in flat]   # 案件名の空白も除いて照合
    pages = txt.count('\f')
    if missing:
        ok(f'[注意] pdftotext で見つからない語: {missing}')
    else:
        ok(f'pdftotext で主要な語をすべて確認（{pages} ページ）')
    if shutil.which('pdffonts'):
        fonts = subprocess.run(['pdffonts', path], capture_output=True, text=True).stdout
        emb = [l for l in fonts.splitlines()[2:] if l.strip()]
        ok('pdffonts: ' + '; '.join(' '.join(l.split()[:1] + l.split()[-5:-2]) for l in emb))


# ---------- 5. はじめにお読みください ----------
def make_readme(path):
    body = f"""ACS株式会社様 デモ用 見本ファイル（{date.today().strftime('%Y-%m-%d')} 作成）

■ すべて架空のデータです
  会社名・人名・案件・数字はデモのために作ったもので、実在の会社・人物・取引とは関係ありません。

■ 落とす順（画面の「内示をここに置く」に入れる）
  1. 内示_前回_2026-08-28.csv       … 前回と同じ → 「変わった点なし」
  2. 今回分は次のどれでも同じ内示が読める（どれか 1 つ、または順番に）
     ・内示_今回_2026-09-25.csv      … CSV
     ・内示_今回_2026-09-25.xlsx     … Excel
     ・内示_今回_2026-09-25.pdf      … 得意先の内示書 PDF（2 ページ）
     ・内示_メール本文_2026-09-25.txt … メール本文（文章で変更を伝えている）
     ・内示_電話メモ_2026-09-25.txt  … 電話で聞いた走り書き（「メモ帳に転記」の欄に貼ってもよい）
  3. 内示_今回_修正版.csv            … 台数・希望日を変えたもの → 表が変わる

■ 新しい案件のデモ（得意先は同じ大和精密。案件を新しく作ってから内示を入れる）
  1. 内示_第3工場_初回_2026-09-30.txt … 電話メモ → この内容で案件を新しく作る（VIS-200 2台、希望 12/10、仕様は標準）
  2. 内示_第3工場_2026-10-03.pdf      … 内示書 PDF（1 ページ）→ 作った案件に入れる（3台、希望 11/28、NG排出シュート）
  3. 内示_第3工場_2026-10-03.csv      … 2 と同じ内容の CSV → 2 の代わりに入れてもよい

■ 今回分の中身（3 件とも同じ変更）
  ・VIS-200 外観検査: 2台 → 3台（3号機追加）、NG排出シュート 3台とも、希望日 11/14（手書きで「できれば 11/7」）
  ・AS-500 組立セル : 2台のまま、希望日 11/10 → 10/15 に前倒し、安全柵 → ライトカーテン仕様（段数は未記載）
  ・CV-CL クリーン搬送: 1台 → 取消（再開未定）

作り直すとき: moc-acs-inventory で  python3 scripts/make-samples.py
"""
    with open(path, 'w', encoding='utf-8') as f:
        f.write(body)
    return True


# ---------- 確認 ----------
def check_pdf(path):
    """pdftotext と pdffonts で、文字が読めるか・フォントが埋め込まれているかを見る"""
    if not shutil.which('pdftotext'):
        ok('pdftotext が無いので文字化けの自動確認はしていません（PDF を開いて目で見てください）')
        return
    # -layout を付けると隣の列の文字が同じ行に混ざるので、素のモードで取り出して空白だけ除く（セル内の折り返しをまたいで照合できる）
    txt = subprocess.run(['pdftotext', path, '-'], capture_output=True, text=True).stdout
    flat = re.sub(r'\s+', '', txt)
    musts = [CUSTOMER, VENDOR, '御中', '内示書', '3号機は増産対応のため追加', 'ライトカーテン仕様', '取消', '別紙']
    missing = [m for m in musts if m not in flat]
    pages = txt.count('\f')
    if missing:
        ok(f'[注意] pdftotext で見つからない語: {missing}')
    else:
        ok(f'pdftotext で主要な語をすべて確認（{pages} ページ）')
    if shutil.which('pdffonts'):
        fonts = subprocess.run(['pdffonts', path], capture_output=True, text=True).stdout
        emb = [l for l in fonts.splitlines()[2:] if l.strip()]
        ok('pdffonts: ' + '; '.join(' '.join(l.split()[:1] + l.split()[-5:-2]) for l in emb))


def main():
    no_copy = '--no-copy' in sys.argv
    new_only = '--new-only' in sys.argv
    os.makedirs(SAMPLE, exist_ok=True)
    made = []
    print('public/sample/ に作る')
    jobs = [
        ('内示_今回_2026-09-25.xlsx', make_xlsx),
        ('内示_今回_2026-09-25.pdf', make_pdf),
        ('内示_メール本文_2026-09-25.txt', make_mail),
        ('内示_電話メモ_2026-09-25.txt', make_memo),
    ]
    new_jobs = [
        ('内示_第3工場_初回_2026-09-30.txt', make_new_memo),
        ('内示_第3工場_2026-10-03.pdf', make_new_pdf),
        ('内示_第3工場_2026-10-03.csv', make_new_csv),
    ]
    for name, fn in (new_jobs if new_only else jobs + new_jobs):
        p = os.path.join(SAMPLE, name)
        if fn(p):
            made.append(name)
            ok(f'{name}  {os.path.getsize(p):,} bytes')
            if name == '内示_今回_2026-09-25.pdf':
                check_pdf(p)
            elif name.endswith('.pdf'):
                check_new_pdf(p)
    if no_copy:
        return
    if new_only:
        print(f'{DOWNLOADS} に新しい案件の分だけ足す（他はそのまま）')
        os.makedirs(DOWNLOADS, exist_ok=True)
        for name in made:
            shutil.copy2(os.path.join(SAMPLE, name), os.path.join(DOWNLOADS, name))
        make_readme(os.path.join(DOWNLOADS, 'はじめにお読みください.txt'))
        for f in sorted(os.listdir(DOWNLOADS)):
            ok(f'{f}  {os.path.getsize(os.path.join(DOWNLOADS, f)):,} bytes')
        return
    print(f'{DOWNLOADS} に並べる（中身は入れ替え）')
    if os.path.isdir(DOWNLOADS):
        for f in os.listdir(DOWNLOADS):
            fp = os.path.join(DOWNLOADS, f)
            (shutil.rmtree if os.path.isdir(fp) else os.remove)(fp)
    os.makedirs(DOWNLOADS, exist_ok=True)
    csvs = sorted(f for f in os.listdir(SAMPLE) if f.endswith('.csv'))
    for name in csvs + made:
        shutil.copy2(os.path.join(SAMPLE, name), os.path.join(DOWNLOADS, name))
    make_readme(os.path.join(DOWNLOADS, 'はじめにお読みください.txt'))
    for f in sorted(os.listdir(DOWNLOADS)):
        ok(f'{f}  {os.path.getsize(os.path.join(DOWNLOADS, f)):,} bytes')


if __name__ == '__main__':
    main()
